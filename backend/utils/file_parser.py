"""
文件解析工具：docx/pdf/xlsx/csv/OCR解析
"""
import os
import io
import csv
import re
import zipfile
import json
import docx
import openpyxl
import requests
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.oxml.ns import qn
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from config import TEXTIN_APP_ID, TEXTIN_SECRET_CODE, DATA_DIR, logger


def iter_block_items(parent):
    """深度解析 Word 底层 XML 树，按物理顺序生成段落和表格"""
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("不支持的节点类型")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def parse_docx_in_order(file_stream) -> str:
    """解析 Docx 为 Markdown"""
    doc = docx.Document(file_stream)
    content = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.replace('\n', ' ').strip()
            if text:
                content.append(text)
        elif isinstance(block, Table):
            content.append("\n")
            for i, row in enumerate(block.rows):
                row_data = [cell.text.replace('\n', ' ').replace('\r', '').strip() for cell in row.cells]
                content.append("| " + " | ".join(row_data) + " |")
                if i == 0:
                    content.append("|" + "|".join(["---"] * len(row.cells)) + "|")
            content.append("\n")
    return "\n".join(content)


def extract_images_from_docx(file_bytes: bytes) -> list:
    """从 Word 文档中提取所有图片"""
    images = []
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            for name in zf.namelist():
                if name.startswith('word/media/') and any(name.lower().endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.bmp', '.gif']):
                    images.append(zf.read(name))
    except Exception as e:
        logger.error(f"提取Word图片失败: {e}")
    return images


def clean_table_html(md_text: str) -> str:
    """
    清理表格中的冗余内容：
    1. 替换 <br> 标签为空格（合并竖向拆分内容）
    2. 横向去重：同一行相邻相同内容只保留第一个
    3. 竖向去重：只对长文本标题（>10字符）去重，签名短文本保留

    风险控制：签名等短文本（<5字符）不做竖向去重，避免误删有效内容
    """
    if not md_text:
        return ""

    # 替换 <br> 标签
    md_text = re.sub(r'<br\s*/?>', ' ', md_text)

    lines = md_text.split('\n')
    result_lines = []
    prev_row_cells = []  # 记录上一行的单元格，用于竖向去重

    for line in lines:
        if not (line.strip().startswith('|') and line.strip().endswith('|')):
            result_lines.append(line)
            continue

        # 表头分割线保持不变
        if re.match(r'^\|[\s\-:|]+\|$', line.strip()):
            result_lines.append(line)
            continue

        # 提取单元格
        cells = [c.strip() for c in line.split('|')[1:-1]]

        # 横向去重：同一行相邻相同只保留第一个
        deduped_cells = []
        prev_cell = ""
        for cell in cells:
            if cell != prev_cell or cell == "":
                deduped_cells.append(cell)
            prev_cell = cell

        # 竖向去重：只对长文本标题去重（>10字符），短文本（签名等）保留
        final_cells = []
        for col_idx, cell in enumerate(deduped_cells):
            if col_idx < len(prev_row_cells):
                prev_cell = prev_row_cells[col_idx]
                # 长文本标题且与上一行相同 → 去重
                # 短文本（签名等）或内容不同 → 保留
                if len(cell) > 10 and cell == prev_cell and cell != "":
                    final_cells.append("")
                else:
                    final_cells.append(cell)
            else:
                final_cells.append(cell)

        # 更新上一行记录（只记录非空单元格）
        new_prev = []
        for col_idx, cell in enumerate(final_cells):
            if col_idx < len(prev_row_cells):
                new_prev.append(cell if cell else prev_row_cells[col_idx])
            else:
                new_prev.append(cell)
        prev_row_cells = new_prev

        # 调整分割线列数
        if result_lines and re.match(r'^\|[\s\-:|]+\|$', result_lines[-1].strip()):
            result_lines[-1] = '|' + '|'.join(['---'] * len(final_cells)) + '|'

        result_lines.append('| ' + ' | '.join(final_cells) + ' |')

    return '\n'.join(result_lines)


def extract_text_via_textin(file_bytes: bytes) -> str:
    """
    调用 TextIn pdf_to_markdown 接口解析文档/图片（scan 模式），
    支持手写签名识别，返回 Markdown 格式内容。
    """
    url = "https://api.textin.com/ai/service/v1/pdf_to_markdown"

    headers = {
        "x-ti-app-id": TEXTIN_APP_ID,
        "x-ti-secret-code": TEXTIN_SECRET_CODE,
        "Content-Type": "application/octet-stream"
    }

    config = {"enable_handwriting": True}

    params = {
        "parse_mode": "scan",  # 纯 OCR 模式，强制识别所有内容包括手写签名
        "table_flavor": "md",
        "apply_document_tree": 1,
        "config": json.dumps(config)
    }

    logger.info(f"TextIn API 调用: parse_mode=scan, enable_handwriting=True")

    try:
        response = requests.post(url, headers=headers, params=params, data=file_bytes, timeout=60)
        result = response.json()

        if result.get("code") == 200:
            raw_markdown = result.get("result", {}).get("markdown", "")
            logger.info(f"TextIn 返回成功，内容长度: {len(raw_markdown)}")
            return clean_table_html(raw_markdown)
        else:
            logger.warning(f"TextIn 接口报错: {result.get('message')}")
    except Exception as e:
        logger.error(f"云端解析异常: {e}")

    return "解析失败"


def read_local_file_content(file_path: str) -> str:
    """读取本地标准文档"""
    if not os.path.exists(file_path):
        return ""

    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".docx":
            with open(file_path, "rb") as f:
                return parse_docx_in_order(f)
        elif ext == ".pdf":
            with open(file_path, "rb") as f:
                return extract_text_via_textin(f.read())
    except Exception as e:
        logger.error(f"解析标准文件 {file_path} 失败: {e}")
    return ""


def get_standard_document_content(major_element: str, matched_rule: str) -> str:
    """读取审核标准原文"""
    if matched_rule:
        for ext in [".docx", ".pdf"]:
            path = os.path.join(DATA_DIR, f"{matched_rule}{ext}")
            content = read_local_file_content(path)
            if content:
                return f"【审核标准原文：{matched_rule}】\n{content}"

    return "未在 data 文件夹下找到对应的标准原文文件，请仅根据通用ISO常识审核。"


def extract_upload_file_text(file_bytes: bytes, filename: str) -> str:
    """
    解析用户上传的证据文件：
    - PDF/图片扫描件：统一使用 TextIn scan 模式 OCR（支持手写签名）
    - Word/Excel/CSV：本地解析
    """
    try:
        ext = filename.lower().split('.')[-1]

        # PDF 和图片扫描件：统一用 TextIn OCR（scan 模式识别手写签名）
        if ext in ['pdf', 'jpg', 'jpeg', 'png', 'bmp', 'gif']:
            logger.info(f"[{filename}] 使用 TextIn scan 模式 OCR 解析（支持手写签名）")
            return extract_text_via_textin(file_bytes)

        # Word 文档
        if ext == 'docx':
            text_content = parse_docx_in_order(io.BytesIO(file_bytes))
            images = extract_images_from_docx(file_bytes)
            for idx, img_bytes in enumerate(images):
                if len(img_bytes) > 1000:
                    ocr_text = extract_text_via_textin(img_bytes)
                    if ocr_text:
                        text_content += f"\n【图片{idx+1} OCR识别内容】\n{ocr_text}"
            return text_content

        if ext == 'doc':
            logger.warning(f"[{filename}] .doc 格式暂不支持，请转换为 .docx")
            return ""

        # Excel 文件
        if ext == 'xlsx':
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            content = []
            for sheet in wb.worksheets:
                content.append(f"\n### {sheet.title}\n")
                for i, r in enumerate(sheet.iter_rows(values_only=True)):
                    row_data = [str(c).replace('\n', ' ').strip() if c is not None else "" for c in r]
                    if not any(row_data):
                        continue
                    content.append("| " + " | ".join(row_data) + " |")
                    if i == 0:
                        content.append("|" + "|".join(["---"] * len(row_data)) + "|")
            return "\n".join(content)

        # CSV 文件
        if ext == 'csv':
            text = file_bytes.decode('utf-8-sig', errors='ignore')
            reader = csv.reader(io.StringIO(text))
            content = []
            for i, row in enumerate(reader):
                row_data = [c.replace('\n', ' ').strip() for c in row]
                if not any(row_data):
                    continue
                content.append("| " + " | ".join(row_data) + " |")
                if i == 0:
                    content.append("|" + "|".join(["---"] * len(row_data)) + "|")
            return "\n".join(content)

    except Exception as e:
        logger.error(f"解析文件异常: {e}")

    return ""